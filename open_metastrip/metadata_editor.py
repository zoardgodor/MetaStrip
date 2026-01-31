"""Module for editing and saving metadata"""
import os
from PIL import Image
import piexif
import mutagen
from PyPDF2 import PdfReader, PdfWriter
from docx import Document
from pptx import Presentation
from openpyxl import load_workbook

SUPPORTED_IMAGE_FORMATS = {'.jpg', '.jpeg', '.tiff', '.tif', '.png', '.bmp', '.gif', '.webp', '.ico', '.svg'}
SUPPORTED_AUDIO_FORMATS = {'.mp3', '.flac', '.ogg', '.wav'}
SUPPORTED_PDF_FORMATS = {'.pdf'}
SUPPORTED_DOCUMENT_FORMATS = {'.docx'}
SUPPORTED_PRESENTATION_FORMATS = {'.pptx'}
SUPPORTED_SPREADSHEET_FORMATS = {'.xlsx'}
SUPPORTED_VIDEO_FORMATS = {'.mp4', '.avi', '.mkv', '.mov', '.flv', '.wmv', '.webm'}


def get_file_extension(file_path):
    return os.path.splitext(file_path)[1].lower()


def save_image_metadata(file_path, metadata_dict):
    try:
        ext = get_file_extension(file_path)
        
        if ext not in {'.jpg', '.jpeg', '.tiff', '.tif'}:
            return False, f"Cannot save EXIF to {ext} files. Only JPEG and TIFF supported."
        
        img = Image.open(file_path)
        
        try:
            exif_data = img.info.get('exif')
            if exif_data:
                exif_dict = piexif.load(exif_data)
            else:
                exif_dict = {
                    "0th": {},
                    "Exif": {},
                    "GPS": {},
                    "1st": {}
                }
        except:
            exif_dict = {
                "0th": {},
                "Exif": {},
                "GPS": {},
                "1st": {}
            }
        
        for key, value in metadata_dict.items():
            if key.startswith("EXIF:"):
                tag_name = key.replace("EXIF:", "")
                try:
                    for ifd_name in ("0th", "Exif", "GPS", "1st"):
                        for tag_id in piexif.TAGS[ifd_name]:
                            if piexif.TAGS[ifd_name][tag_id]["name"] == tag_name:
                                tag_info = piexif.TAGS[ifd_name][tag_id]
                                tag_type = tag_info.get("type", 2)  
                                
                                if tag_type == 3: 
                                    try:
                                        exif_dict[ifd_name][tag_id] = int(value)
                                    except:
                                        exif_dict[ifd_name][tag_id] = 0
                                elif tag_type == 4:  
                                    try:
                                        exif_dict[ifd_name][tag_id] = int(value)
                                    except:
                                        exif_dict[ifd_name][tag_id] = 0
                                elif tag_type == 5 or tag_type == 10:  
                                    try:
                                        float_val = float(value)
                                        exif_dict[ifd_name][tag_id] = (int(float_val * 1000), 1000)
                                    except:
                                        exif_dict[ifd_name][tag_id] = (0, 1)
                                else:  
                                    exif_dict[ifd_name][tag_id] = str(value).encode('utf-8')
                                break
                except:
                    pass
        
        exif_bytes = piexif.dump(exif_dict)
        img.save(file_path, "jpeg" if ext in {'.jpg', '.jpeg'} else "tiff", exif=exif_bytes)
        return True, "Image metadata saved successfully"
        
    except Exception as e:
        return False, f"Error saving image metadata: {str(e)}"


def save_audio_metadata(file_path, metadata_dict):
    try:
        audio = mutagen.File(file_path, easy=True)
        if audio is None:
            audio = mutagen.File(file_path)
            if audio is None:
                return False, "Cannot read audio file"
        
        for key, value in metadata_dict.items():
            if key.startswith("ID3:"):
                tag_name = key.replace("ID3:", "")
                audio[tag_name] = str(value)
        
        audio.save()
        return True, "Audio metadata saved successfully"
        
    except Exception as e:
        return False, f"Error saving audio metadata: {str(e)}"


def save_pdf_metadata(file_path, metadata_dict):
    try:
        reader = PdfReader(file_path)
        writer = PdfWriter()
        
        for page in reader.pages:
            writer.add_page(page)
        
        pdf_meta = {}
        for key, value in metadata_dict.items():
            if key.startswith("PDF:"):
                tag_name = key.replace("PDF:", "")
                pdf_meta[f"/{tag_name}"] = str(value)
        
        if pdf_meta:
            writer.add_metadata(pdf_meta)
        
        with open(file_path, 'wb') as f:
            writer.write(f)
        
        return True, "PDF metadata saved successfully"
        
    except Exception as e:
        return False, f"Error saving PDF metadata: {str(e)}"


def save_metadata(file_path, metadata_dict):
    """Save metadata based on file type"""
    ext = get_file_extension(file_path)
    
    if ext in SUPPORTED_IMAGE_FORMATS:
        return save_image_metadata(file_path, metadata_dict)
    elif ext in SUPPORTED_AUDIO_FORMATS:
        return save_audio_metadata(file_path, metadata_dict)
    elif ext in SUPPORTED_PDF_FORMATS:
        return save_pdf_metadata(file_path, metadata_dict)
    elif ext in SUPPORTED_DOCUMENT_FORMATS:
        return save_docx_metadata(file_path, metadata_dict)
    elif ext in SUPPORTED_PRESENTATION_FORMATS:
        return save_pptx_metadata(file_path, metadata_dict)
    elif ext in SUPPORTED_SPREADSHEET_FORMATS:
        return save_xlsx_metadata(file_path, metadata_dict)
    else:
        return False, f"Unsupported file format: {ext}"


def save_docx_metadata(file_path, metadata_dict):
    try:
        doc = Document(file_path)
        props = doc.core_properties
        
        for key, value in metadata_dict.items():
            if key == "DOCX:Cím":
                props.title = str(value)
            elif key == "DOCX:Tárgy":
                props.subject = str(value)
            elif key == "DOCX:Szerző":
                props.author = str(value)
            elif key == "DOCX:Kulcsszavak":
                props.keywords = str(value)
            elif key == "DOCX:Megjegyzések":
                props.comments = str(value)
        
        doc.save(file_path)
        return True, "DOCX metaadatok mentve"
    except Exception as e:
        return False, f"Hiba a DOCX metaadatok mentésekor: {str(e)}"


def save_pptx_metadata(file_path, metadata_dict):
    try:
        prs = Presentation(file_path)
        props = prs.core_properties
        
        for key, value in metadata_dict.items():
            if key == "PPTX:Cím":
                props.title = str(value)
            elif key == "PPTX:Tárgy":
                props.subject = str(value)
            elif key == "PPTX:Szerző":
                props.author = str(value)
            elif key == "PPTX:Kulcsszavak":
                props.keywords = str(value)
            elif key == "PPTX:Megjegyzések":
                props.comments = str(value)
        
        prs.save(file_path)
        return True, "PPTX metaadatok mentve"
    except Exception as e:
        return False, f"Hiba a PPTX metaadatok mentésekor: {str(e)}"


def save_xlsx_metadata(file_path, metadata_dict):
    try:
        wb = load_workbook(file_path)
        props = wb.properties
        
        for key, value in metadata_dict.items():
            if key == "XLSX:Cím":
                props.title = str(value)
            elif key == "XLSX:Tárgy":
                props.subject = str(value)
            elif key == "XLSX:Szerző":
                props.author = str(value)
            elif key == "XLSX:Kulcsszavak":
                props.keywords = str(value)
            elif key == "XLSX:Megjegyzések":
                props.comments = str(value)
        
        wb.save(file_path)
        return True, "XLSX metaadatok mentve"
    except Exception as e:
        return False, f"Hiba az XLSX metaadatok mentésekor: {str(e)}"
